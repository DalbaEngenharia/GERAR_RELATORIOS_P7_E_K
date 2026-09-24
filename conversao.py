from pathlib import Path
import xml.etree.ElementTree as ET
from openpyxl import Workbook


def converter_xmls_para_xlsx():
    pasta = Path(r"C:\Users\DALBAPY\Desktop\Arquivos_Protheus")

    if not pasta.exists():
        print(f"Pasta não encontrada: {pasta}")
        return

    arquivos_xml = list(pasta.glob("*.xml"))

    if not arquivos_xml:
        print("Nenhum arquivo XML encontrado.")
        return

    ns = {
        "ss": "urn:schemas-microsoft-com:office:spreadsheet"
    }

    print(f"Encontrados {len(arquivos_xml)} XML(s).\n")

    for arquivo_xml in arquivos_xml:

        try:
            # Lê o XML
            tree = ET.parse(arquivo_xml)
            root = tree.getroot()

            # Cria o XLSX
            wb = Workbook()

            # Remove a aba padrão
            wb.remove(wb.active)

            # Localiza as planilhas do XML
            worksheets = root.findall(".//ss:Worksheet", ns)

            if not worksheets:
                raise ValueError("Nenhuma Worksheet encontrada no XML.")

            # Percorre as abas
            for indice, worksheet in enumerate(worksheets, start=1):

                nome_aba = worksheet.get(
                    "{urn:schemas-microsoft-com:office:spreadsheet}Name"
                )

                if not nome_aba:
                    nome_aba = f"Planilha{indice}"

                # Excel permite no máximo 31 caracteres
                nome_aba = nome_aba[:31]

                ws = wb.create_sheet(title=nome_aba)

                # Localiza a tabela
                table = worksheet.find("ss:Table", ns)

                if table is None:
                    continue

                linha = 1

                # Percorre as linhas
                for row in table.findall("ss:Row", ns):

                    coluna = 1

                    # Índice explícito da linha
                    index_linha = row.get(
                        "{urn:schemas-microsoft-com:office:spreadsheet}Index"
                    )

                    if index_linha:
                        linha = int(index_linha)

                    # Percorre as células
                    for cell in row.findall("ss:Cell", ns):

                        # Índice explícito da coluna
                        index_coluna = cell.get(
                            "{urn:schemas-microsoft-com:office:spreadsheet}Index"
                        )

                        if index_coluna:
                            coluna = int(index_coluna)

                        data = cell.find("ss:Data", ns)

                        valor = None

                        if data is not None:
                            valor = data.text

                            tipo = data.get(
                                "{urn:schemas-microsoft-com:office:spreadsheet}Type"
                            )

                            # Converte números
                            if tipo == "Number" and valor:
                                try:
                                    numero = float(valor)

                                    if numero.is_integer():
                                        valor = int(numero)
                                    else:
                                        valor = numero

                                except ValueError:
                                    pass

                            # Converte booleanos
                            elif tipo == "Boolean":
                                valor = valor == "1"

                        ws.cell(
                            row=linha,
                            column=coluna,
                            value=valor
                        )

                        coluna += 1

                    linha += 1

            # Nome do XLSX igual ao XML
            arquivo_xlsx = arquivo_xml.with_suffix(".xlsx")

            # Salva o XLSX
            wb.save(arquivo_xlsx)

            print(
                f"OK: {arquivo_xml.name} -> {arquivo_xlsx.name}"
            )

            # Remove o XML somente após salvar o XLSX com sucesso
            arquivo_xml.unlink()

            print(
                f"XML removido: {arquivo_xml.name}\n"
            )

        except Exception as erro:

            print(
                f"ERRO: {arquivo_xml.name} -> {erro}\n"
            )


